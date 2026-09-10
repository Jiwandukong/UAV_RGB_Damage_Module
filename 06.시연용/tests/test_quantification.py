"""Demo table and image-bundle contract; no real OBJ, checkpoint or GPU."""

from copy import deepcopy
import csv
import json
from pathlib import Path
import re
import sys
from types import SimpleNamespace

import numpy as np
from openpyxl import load_workbook
from PIL import Image
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "code"))
from demo512 import quantification as quant
from demo512.data import CLASSES, OVERLAY_ALPHA, render_overlay, sha256_file
from demo512.report_contract import SCHEMA_PATH


EXPECTED_COLUMNS = [
    "image", "damage_id", "damage_type", "damage_name_ko", "pixel_nodes_json",
    "world_center_x_m", "world_center_y_m", "world_center_z_m",
    "length_px", "length_m", "width_px", "width_m", "area_m2",
    "source_image_path", "tile_original_paths_json", "tile_overlay_paths_json",
]


def inputs(label="CRC"):
    source = {"id": "fixture", "source_filename": "fixture.JPG"}
    annotation = {
        "damage_id": f"fixture__{label}_0001", "shape_index": 0, "label": label,
        "tiles": [{"tile_id": "left"}, {"tile_id": "right"}],
    }
    tiles = {
        name: {"id": name, "source_id": "fixture", "x0": x0, "y0": 0,
               "valid_width": width, "valid_height": 4,
               "damage_ids": [annotation["damage_id"]]}
        for name, x0, width in (("left", 0, 512), ("right", 512, 1))
    }
    spatial = {
        "xyz_valid": True, "measurement_valid": True,
        "world_x_m": 243001.25, "world_y_m": 431002.5, "world_z_m": 77.75,
        "image_pixel_x": 512.0, "image_pixel_y": 1.0,
        "mesh_ray_t_m": 1.0,
        "nodes_image_xy_json": "[[510.5,0.5],[512.5,1.5]]",
        "gsd_length_m_per_px": 0.001, "gsd_width_m_per_px": 0.001,
        "length_px": 2.0 if label == "CRC" else None,
        "width_px": 1.0 if label == "CRC" else None,
        "length_m": 0.002 if label == "CRC" else None,
        "width_m": 0.001 if label == "CRC" else None,
        "area_m2": None if label == "CRC" else 0.000002,
    }
    return source, annotation, spatial, tiles


def test_exact_existing_schema_and_crc_min_rectangle_values_are_preserved():
    source, annotation, spatial, tiles = inputs()
    # Deliberately noninteger/curved-line rectangle values: never replace these
    # with line-path length, force width=1, or claim physical opening width.
    spatial.update(length_px=37.625, width_px=11.875,
                   length_m=0.037625, width_m=0.011875)
    before = deepcopy((source, annotation, spatial, tiles))
    row, mapping = quant.build_row(source, annotation, spatial, "D000001", tiles, 1000.0)
    assert list(row) == EXPECTED_COLUMNS
    schema = json.loads(SCHEMA_PATH.read_text())
    assert schema["additionalProperties"] is False
    assert schema["required"] == EXPECTED_COLUMNS
    assert set(schema["properties"]) == set(row)
    assert re.fullmatch(schema["properties"]["damage_id"]["pattern"], row["damage_id"])
    for name, value in row.items():
        property_schema = schema["properties"][name]
        if "enum" in property_schema:
            assert value in property_schema["enum"]
        else:
            types = property_schema["type"]
            types = [types] if isinstance(types, str) else types
            assert ((value is None and "null" in types)
                    or (isinstance(value, str) and "string" in types)
                    or (isinstance(value, (int, float)) and "number" in types))
    for name in ("length_px", "width_px", "length_m", "width_m"):
        assert row[name] == spatial[name]
    assert row["area_m2"] is None
    assert mapping["source_annotation_id"] == annotation["damage_id"]
    assert mapping["raw_legacy_mapping_and_measurement"] == spatial
    assert (source, annotation, spatial, tiles) == before


@pytest.mark.parametrize("label, korean", [("DLM", "박리"), ("SPL", "박락")])
def test_area_classes_leave_length_and_width_blank(label, korean):
    source, annotation, spatial, tiles = inputs(label)
    row, mapping = quant.build_row(source, annotation, spatial, "D000001", tiles, 1000.0)
    assert row["damage_type"] == label and row["damage_name_ko"] == korean
    assert all(row[key] is None for key in ("length_px", "length_m", "width_px", "width_m"))
    assert row["area_m2"] == spatial["area_m2"]
    assert mapping["review"]["physical_values_exported"] is True


def test_one_annotation_crossing_tiles_has_one_row_and_representative_tile_first():
    source, annotation, spatial, tiles = inputs()
    row, mapping = quant.build_row(source, annotation, spatial, "D000007", tiles, 1000.0)
    assert isinstance(row, dict) and row["damage_id"] == "D000007"
    original = json.loads(row["tile_original_paths_json"])
    overlays = json.loads(row["tile_overlay_paths_json"])
    assert original == ["512원본타일/right.png", "512원본타일/left.png"]
    assert overlays == ["512라벨오버레이/right.png", "512라벨오버레이/left.png"]
    assert [Path(path).name for path in original] == [Path(path).name for path in overlays]
    assert mapping["tile_ids"] == ["right", "left"]
    assert annotation["tiles"] == [{"tile_id": "left"}, {"tile_id": "right"}]
    assert all(not Path(path).is_absolute() and ".." not in Path(path).parts
               for path in original + overlays + [row["source_image_path"]])


def test_no_hit_keeps_pixel_geometry_links_and_raw_record_but_nulls_world_and_metres():
    source, annotation, spatial, tiles = inputs()
    spatial.update(xyz_valid=False, measurement_valid=False,
                   xyz_miss_reason="mesh_ray_no_hit", measurement_miss_reason="missing_stencil")
    before = deepcopy(spatial)
    row, mapping = quant.build_row(source, annotation, spatial, "D000001", tiles, 1000.0)
    assert all(row[key] is None for key in (
        "world_center_x_m", "world_center_y_m", "world_center_z_m", "length_m", "width_m", "area_m2"))
    assert row["length_px"] == 2.0 and row["width_px"] == 1.0
    assert row["pixel_nodes_json"] == spatial["nodes_image_xy_json"]
    assert len(json.loads(row["tile_original_paths_json"])) == 2
    assert mapping["review"]["reasons"] == ["mesh_ray_no_hit", "missing_stencil"]
    assert mapping["review"]["physical_values_exported"] is False
    assert mapping["raw_legacy_mapping_and_measurement"] == before == spatial


@pytest.mark.parametrize("ratio, allowed", [(10.0, True), (10.001, False)])
def test_frontal_review_threshold_preserves_raw_legacy_values(ratio, allowed):
    source, annotation, spatial, tiles = inputs()
    spatial.update(gsd_length_m_per_px=ratio * 0.001,
                   length_m=spatial["length_px"] * ratio * 0.001)
    before = deepcopy(spatial)
    review = quant.review_measurement(spatial, 1000.0)
    row, mapping = quant.build_row(source, annotation, spatial, "D000001", tiles, 1000.0)
    assert review["physical_values_exported"] is allowed
    assert review["max_directional_gsd_to_frontal_ratio"] == pytest.approx(ratio)
    assert review["frontal_reference_m_per_px"] == 0.001
    assert row["world_center_x_m"] == spatial["world_x_m"]
    assert row["length_px"] == spatial["length_px"]
    if allowed:
        assert row["length_m"] == spatial["length_m"]
    else:
        assert all(row[key] is None for key in ("length_m", "width_m", "area_m2"))
        assert review["reasons"] == ["local_gsd_exceeds_10x_frontal_reference_manual_review"]
    assert mapping["raw_legacy_mapping_and_measurement"] == before == spatial


@pytest.mark.parametrize("kind", ["empty", "duplicate", "wrong_annotation", "wrong_source"])
def test_invalid_annotation_tile_link_is_rejected(kind):
    source, annotation, spatial, tiles = inputs()
    if kind == "empty":
        annotation["tiles"] = []
    elif kind == "duplicate":
        annotation["tiles"] = [{"tile_id": "left"}, {"tile_id": "left"}]
    elif kind == "wrong_annotation":
        tiles["left"]["damage_ids"] = ["different_annotation"]
    else:
        tiles["left"]["source_id"] = "different_photo"
    with pytest.raises(ValueError):
        quant.build_row(source, annotation, spatial, "D000001", tiles, 1000.0)


def make_export_fixture(tmp_path, monkeypatch):
    root = tmp_path / "data"
    root.mkdir()
    source, annotation, spatial, tile_lookup = inputs()
    image = Image.new("RGB", (513, 4), (60, 80, 100))
    image.save(root / "fixture.JPG")
    shape = {"label": "CRC", "shape_type": "linestrip", "points": [[511, 1], [512, 1]]}
    (root / "fixture.json").write_text(json.dumps({"shapes": [shape]}), encoding="utf-8")
    annotation["original_shape"] = shape
    source.update(width=513, height=4, image_path="fixture.JPG", label_path="fixture.json",
                  image_sha256=sha256_file(root / "fixture.JPG"),
                  label_sha256=sha256_file(root / "fixture.json"), annotations=[annotation])
    with Image.open(root / "fixture.JPG") as opened:
        decoded = opened.convert("RGB")
    for tile in tile_lookup.values():
        rgb = decoded.crop((tile["x0"], 0, tile["x0"] + 512, 512))
        tile["image_path"] = f"{tile['id']}.png"
        rgb.save(root / tile["image_path"])
        masks = {label: Image.new("L", (512, 512), 0) for label in CLASSES}
        masks["CRC"].putpixel((511 if tile["id"] == "left" else 0, 1), 255)
        tile["mask_paths"] = {}
        for label, mask in masks.items():
            relative = f"{tile['id']}_{label}.png"
            mask.save(root / relative)
            tile["mask_paths"][label] = relative
        tile["positive_pixels"] = {label: int(np.count_nonzero(mask)) for label, mask in masks.items()}
        # An opaque historical preview must not be copied into the new output.
        tile["overlay_path"] = f"{tile['id']}_old_opaque.png"
        render_overlay(rgb, masks, alpha=1).save(root / tile["overlay_path"])
    dataset = {"tile_size": 512, "classes": list(CLASSES), "sources": [source],
               "tiles": list(tile_lookup.values())}
    (root / "dataset.json").write_text(json.dumps(dataset), encoding="utf-8")
    mesh = tmp_path / "fixture.obj"
    mesh.write_text("mock mesh; never loaded", encoding="utf-8")
    contract = SimpleNamespace(to_dict=lambda: {"fixture_contract": True})
    surface = SimpleNamespace(vertex_count=3, face_count=1, bounds=[[0, 0, 0], [1, 1, 0]],
                              ray_backend="fixture_cpu")
    context = SimpleNamespace(intrinsics=SimpleNamespace(focal_length_px=1000.0),
                              to_dict=lambda: {"fixture_context": True})
    calls = []
    monkeypatch.setattr(quant, "read_dji_xmp", lambda path: {key: "1" for key in quant.REQUIRED_XMP})
    monkeypatch.setattr(quant, "load_mesh_asset_contract", lambda path: contract)
    monkeypatch.setattr(quant, "verify_mesh_file", lambda *args: {"verified": True})
    monkeypatch.setattr(quant, "MeshSurfaceIndex", lambda *args, **kwargs: surface)
    monkeypatch.setattr(quant, "verify_loaded_mesh_geometry", lambda **kwargs: {"verified": True})

    def build_context(image_path, actual_surface, **kwargs):
        assert actual_surface is surface
        assert kwargs["mrk_path"] is None and kwargs["representative_mode"] == "centroid"
        calls.append(image_path)
        return context

    def measure(data_root, actual_source, actual_context):
        assert data_root == root and actual_context is context
        assert actual_source == source
        return [{"annotation": actual_source["annotations"][0], "spatial": deepcopy(spatial),
                 "base": {"area_px": 2}}]

    monkeypatch.setattr(quant, "build_mesh_ray_context_with_surface", build_context)
    monkeypatch.setattr(quant, "quantify_source", measure)
    return root, mesh, calls


def test_real_csv_xlsx_png_export_is_gt_only_self_contained_and_preserves_input(tmp_path, monkeypatch):
    root, mesh, calls = make_export_fixture(tmp_path, monkeypatch)
    before = {path.relative_to(root): path.read_bytes() for path in root.rglob("*") if path.is_file()}
    output = tmp_path / "new_demo_output"
    report = quant.export_demo(root, output, mesh)
    assert report["rows"] == 1 and report["images"] == 1 and report["tiles"] == 2
    assert report["mapped_rows"] == report["physical_values_exported"] == 1
    assert len(calls) == 1
    assert (output / "damage_results.csv").read_bytes().startswith(b"\xef\xbb\xbf")
    with (output / "damage_results.csv").open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        assert reader.fieldnames == EXPECTED_COLUMNS
        rows = list(reader)
    assert len(rows) == 1
    row = rows[0]
    workbook = load_workbook(output / "damage_results.xlsx", read_only=True)
    try:
        assert workbook.sheetnames == ["Damage_Details"]
        cells = list(workbook.active.values)
        assert list(cells[0]) == EXPECTED_COLUMNS and len(cells) == 2
        assert cells[1][12] is None
    finally:
        workbook.close()
    assert (output / row["source_image_path"]).read_bytes() == (root / "fixture.JPG").read_bytes()
    original_paths = json.loads(row["tile_original_paths_json"])
    overlay_paths = json.loads(row["tile_overlay_paths_json"])
    assert Path(original_paths[0]).stem == "right"
    for original, overlay in zip(original_paths, overlay_paths, strict=True):
        assert Path(original).name == Path(overlay).name
        with Image.open(output / original) as rgb, Image.open(output / overlay) as painted:
            assert rgb.size == painted.size == (512, 512)
            assert rgb.mode == painted.mode == "RGB"
            x = 0 if Path(original).stem == "right" else 511
            expected = tuple(int(0.5 * value + 0.5 * color)
                             for value, color in zip(rgb.getpixel((x, 1)), (0, 255, 0)))
            assert painted.getpixel((x, 1)) == expected
            assert painted.getpixel((x, 2)) == rgb.getpixel((x, 2))
            assert painted.getpixel((x, 5)) == (0, 0, 0)
    record = json.loads((output / "정량계산기록.json").read_text())
    connections = json.loads((output / "연결정보.json").read_text())
    assert record["result_source"] == "human_reviewed_labels_not_model_predictions"
    assert record["model_loaded"] is False and record["source_resized"] is False
    assert record["crc_width_is_physical_aperture"] is False
    assert record["columns"] == connections["columns"] == EXPECTED_COLUMNS
    assert record["damages"][0]["raw_legacy_mapping_and_measurement"]["length_px"] == 2.0
    assert connections["overlay"]["alpha"] == OVERLAY_ALPHA == 0.5
    assert len(connections["damages"]) == 1 and len(connections["tiles"]) == 2
    assert {path.relative_to(root): path.read_bytes() for path in root.rglob("*") if path.is_file()} == before
    assert not list(tmp_path.glob(".demo-quantification-*"))


def test_export_refuses_existing_destination_before_reading_assets(tmp_path):
    root, output = tmp_path / "data", tmp_path / "existing"
    root.mkdir()
    output.mkdir()
    sentinel = output / "keep.txt"
    sentinel.write_text("user-owned output", encoding="utf-8")
    with pytest.raises(FileExistsError, match="overwrite"):
        quant.export_demo(root, output, tmp_path / "absent.obj")
    assert sentinel.read_text() == "user-owned output"
    assert sorted(path.name for path in output.iterdir()) == ["keep.txt"]


def test_export_refuses_output_inside_immutable_training_data(tmp_path):
    root = tmp_path / "data"
    root.mkdir()
    with pytest.raises(ValueError, match="immutable training dataset"):
        quant.export_demo(root, root / "new-output", tmp_path / "absent.obj")
    assert list(root.iterdir()) == []


def test_changed_source_is_rejected_before_mesh_load_or_output_creation(tmp_path, monkeypatch):
    root, mesh, calls = make_export_fixture(tmp_path, monkeypatch)
    (root / "fixture.json").write_text('{"shapes": []}', encoding="utf-8")
    monkeypatch.setattr(quant, "MeshSurfaceIndex", lambda *args, **kwargs: pytest.fail("mesh must not load"))
    output = tmp_path / "new-output"
    with pytest.raises(ValueError, match="source image or labels changed"):
        quant.export_demo(root, output, mesh)
    assert not output.exists() and not calls
