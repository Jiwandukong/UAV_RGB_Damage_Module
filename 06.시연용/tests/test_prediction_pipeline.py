"""Raw-photo outputs must contain actual predicted masks, never supplied GT.

Only the heavyweight model and mesh loader are faked. Native pixel geometry,
CSV/Excel writers, tile saving, linking, and atomic publication run for real.
"""

import builtins
import csv
import hashlib
import json
from pathlib import Path
import sys
from types import SimpleNamespace

import numpy as np
from openpyxl import load_workbook
from PIL import Image
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "code"))
from demo512 import prediction_pipeline as pipeline
from demo512.data import CLASSES, OVERLAY_ALPHA, render_overlay


EXPECTED_COLUMNS = [
    "image", "damage_id", "damage_type", "damage_name_ko", "pixel_nodes_json",
    "world_center_x_m", "world_center_y_m", "world_center_z_m",
    "length_px", "length_m", "width_px", "width_m", "area_m2",
    "source_image_path", "tile_original_paths_json", "tile_overlay_paths_json",
]
WORLD_COLUMNS = ["world_center_x_m", "world_center_y_m", "world_center_z_m"]
PHYSICAL_COLUMNS = ["length_m", "width_m", "area_m2"]


def source_hashes(folder):
    return {str(path.relative_to(folder)): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in folder.rglob("*") if path.is_file()}


def output_rows(folder):
    with (folder / "damage_results.csv").open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        assert reader.fieldnames == EXPECTED_COLUMNS
        return list(reader)


@pytest.fixture
def fixture(tmp_path, monkeypatch):
    images = tmp_path / "원본 사진"
    images.mkdir()
    image = images / "unlabelled.png"
    pixels = np.zeros((8, 519, 3), dtype=np.uint8)
    pixels[:] = (60, 80, 100)
    pixels[:, :, 0] = np.arange(519, dtype=np.uint16) % 256
    Image.fromarray(pixels).save(image)
    # Deliberately unusable labels: no raw-image path may consult either file.
    (images / "unlabelled.json").write_text("THIS IS NOT LABEL DATA", encoding="utf-8")
    (images / "dataset.json").write_text("THIS IS NOT A DATASET", encoding="utf-8")
    record = tmp_path / "学習記録.json"
    record.write_text('{"fixture": true}', encoding="utf-8")
    mesh = tmp_path / "existing.obj"
    mesh.write_text("mock OBJ; never loaded", encoding="utf-8")
    output = tmp_path / "結果"
    state = {"images": images, "image": image, "record": record, "mesh": mesh,
             "output": output, "pixels": pixels, "empty": False,
             "fail_after_tile": False, "ratio": 1.0, "engine_calls": [],
             "predictions": [], "mesh_loads": [], "context_calls": []}

    class FakeEngine:
        def __init__(self, model_record, device="auto", threshold=0.5):
            state["engine_calls"].append((Path(model_record), device, threshold))
            assert Path(model_record).resolve() == record
            self.metadata = {
                "prediction_source": "actual_model_semantic_logits",
                "ground_truth_used_for_prediction": False,
                "ground_truth_used_for_tile_selection": False,
                "dataset_manifest_read": False,
                "input_size": [512, 512], "source_resized": False,
                "input_resized": False, "device": device,
            }

        def predict(self, rgb, *, on_tile=None):
            assert rgb.dtype == np.uint8 and rgb.shape[2] == 3
            height, width = rgb.shape[:2]
            masks = {key: np.zeros((height, width), dtype=bool) for key in (1, 2, 3)}
            if not state["empty"] and rgb.any():
                masks[1][1:3, 509:518] = True  # One component crosses x=512.
                masks[2][2:5, 1:5] = True
                masks[3][3:6, 3:7] = True  # DLM/SPL predictions overlap.
            state["predictions"].append({"rgb": rgb.copy(), "masks": masks})
            tiles = []
            for y0 in range(0, height, 512):
                for x0 in range(0, width, 512):
                    valid_width, valid_height = min(512, width - x0), min(512, height - y0)
                    tile = {"id_suffix": f"x{x0:05d}_y{y0:05d}", "x0": x0, "y0": y0,
                            "valid_width": valid_width, "valid_height": valid_height}
                    rgb512 = np.zeros((512, 512, 3), dtype=np.uint8)
                    rgb512[:valid_height, :valid_width] = rgb[y0:y0 + valid_height, x0:x0 + valid_width]
                    tile_masks = {}
                    for key, label in enumerate(CLASSES, 1):
                        tile_masks[label] = np.zeros((512, 512), dtype=bool)
                        tile_masks[label][:valid_height, :valid_width] = masks[key][
                            y0:y0 + valid_height, x0:x0 + valid_width]
                    tiles.append(tile)
                    if on_tile is not None:
                        on_tile(tile, rgb512, tile_masks)
                    if state["fail_after_tile"]:
                        raise RuntimeError("fixture inference failed after first saved tile")
            return {"class_masks": masks, "tiles": tiles, "elapsed_seconds": 0.125}

    contract = SimpleNamespace(to_dict=lambda: {"fixture_contract": True})
    surface = SimpleNamespace(vertex_count=3, face_count=1,
                              bounds=[[0, 0, 0], [1, 1, 0]], ray_backend="fixture_cpu")
    context = SimpleNamespace(
        intrinsics=SimpleNamespace(focal_length_px=1000.0, width=519, height=8),
        to_dict=lambda: {"fixture_context": True},
    )

    def map_region(base, **kwargs):
        is_crack = base["class_id"] == 1
        return {
            **base, "xyz_valid": True, "measurement_valid": True,
            "world_x_m": 243001.0, "world_y_m": 431001.0, "world_z_m": 70.0,
            "image_pixel_x": base["centroid_x_px"], "image_pixel_y": base["centroid_y_px"],
            "mesh_ray_t_m": 1.0,
            "gsd_length_m_per_px": 0.001 * state["ratio"],
            "gsd_width_m_per_px": 0.001,
            "length_m": base["length_px"] * 0.001 * state["ratio"] if is_crack else None,
            "width_m": base["width_px"] * 0.001 if is_crack else None,
            "area_m2": None if is_crack else base["area_px"] * 0.000001,
            "nodes_image_xy_json": "[]",
        }

    context.image_region_to_world3d = map_region

    def mesh_loader(path, **kwargs):
        assert Path(path) == mesh
        state["mesh_loads"].append(kwargs)
        return surface

    def build_context(image_path, actual_surface, **kwargs):
        assert actual_surface is surface
        assert kwargs["mrk_path"] is None
        state["context_calls"].append(Path(image_path))
        return context

    monkeypatch.setattr(pipeline, "RawInferenceEngine", FakeEngine)
    monkeypatch.setattr(pipeline, "load_mesh_asset_contract", lambda path: contract)
    monkeypatch.setattr(pipeline, "verify_mesh_file", lambda *args: {"verified": True})
    monkeypatch.setattr(pipeline, "MeshSurfaceIndex", mesh_loader)
    monkeypatch.setattr(pipeline, "verify_loaded_mesh_geometry", lambda **kwargs: {"verified": True})
    monkeypatch.setattr(pipeline, "build_mesh_ray_context_with_surface", build_context)
    monkeypatch.setattr(pipeline, "read_dji_xmp", lambda path: {})
    return state


def export(fixture, **kwargs):
    options = dict(images=fixture["images"], output=fixture["output"],
                   model_record=fixture["record"], mesh=fixture["mesh"], device="cpu",
                   ray_backend="warp", warp_device="cpu")
    options.update(kwargs)
    return pipeline.export_predictions(**options)


def enable_metadata(monkeypatch):
    metadata = {
        "GpsLatitude": "36.5", "GpsLongitude": "127.5", "AbsoluteAltitude": "80",
        "GimbalYawDegree": "0", "GimbalPitchDegree": "0", "GimbalRollDegree": "0",
        "CalibratedFocalLength": "1000", "CalibratedOpticalCenterX": "259.5",
        "CalibratedOpticalCenterY": "4",
    }
    monkeypatch.setattr(pipeline, "read_dji_xmp", lambda path: dict(metadata))


def test_raw_photo_exports_exact_16_columns_and_only_predicted_overlay_links(fixture):
    before = source_hashes(fixture["images"])
    result = export(fixture)
    output = fixture["output"]
    rows = output_rows(output)
    assert result["images"] == 1 and result["tiles"] == 2 and result["rows"] == 3
    assert {row["damage_type"] for row in rows} == {"CRC", "DLM", "SPL"}
    assert [row["damage_id"] for row in rows] == ["D000001", "D000002", "D000003"]
    assert all(row[key] == "" for row in rows for key in WORLD_COLUMNS + PHYSICAL_COLUMNS)
    assert result["mapped_rows"] == 0 and result["physical_values_exported"] == 0
    assert result["physical_values_withheld"] == 3
    assert (output / "damage_results.csv").read_bytes().startswith(b"\xef\xbb\xbf")
    assert not (output / "512라벨오버레이").exists()
    for row in rows:
        original = json.loads(row["tile_original_paths_json"])
        overlays = json.loads(row["tile_overlay_paths_json"])
        assert original and len(original) == len(overlays)
        assert [Path(path).name for path in original] == [Path(path).name for path in overlays]
        for relative in [row["source_image_path"], *original, *overlays]:
            assert not Path(relative).is_absolute() and ".." not in Path(relative).parts
            assert (output / relative).is_file()
        assert all(Path(path).parts[0] == "512모델예측오버레이" for path in overlays)
        assert json.loads(row["pixel_nodes_json"])
    crc = next(row for row in rows if row["damage_type"] == "CRC")
    assert len(json.loads(crc["tile_original_paths_json"])) == 2
    assert float(crc["length_px"]) > 0 and float(crc["width_px"]) > 0
    assert source_hashes(fixture["images"]) == before
    assert (output / "원본사진/unlabelled.png").read_bytes() == fixture["image"].read_bytes()
    assert len(fixture["mesh_loads"]) == 1 and len(fixture["engine_calls"]) == 1
    assert not fixture["context_calls"]
    workbook = load_workbook(output / "damage_results.xlsx", read_only=True)
    try:
        assert workbook.sheetnames == ["Damage_Details"]
        values = list(workbook.active.values)
        assert list(values[0]) == EXPECTED_COLUMNS
        assert len(values) == 4
        for row, cells in zip(rows, values[1:]):
            for column, cell in zip(EXPECTED_COLUMNS, cells):
                if cell is None:
                    assert row[column] == ""
                elif isinstance(cell, (int, float)):
                    assert float(row[column]) == cell
                else:
                    assert cell == row[column]
    finally:
        workbook.close()


def test_native_tile_pixels_alpha_and_independent_full_image_masks(fixture):
    export(fixture)
    output = fixture["output"]
    prediction = fixture["predictions"][0]
    originals = sorted((output / "512원본타일").glob("*.png"))
    assert len(originals) == 2
    for path in originals:
        x0 = 512 if "x00512" in path.stem else 0
        valid_width = min(512, 519 - x0)
        expected_rgb = np.zeros((512, 512, 3), dtype=np.uint8)
        expected_rgb[:8, :valid_width] = fixture["pixels"][:, x0:x0 + valid_width]
        with Image.open(path) as opened:
            assert opened.mode == "RGB" and opened.size == (512, 512)
            assert np.array_equal(np.asarray(opened), expected_rgb)
        tile_masks = {}
        for key, label in enumerate(CLASSES, 1):
            mask = np.zeros((512, 512), dtype=np.uint8)
            mask[:8, :valid_width] = prediction["masks"][key][:, x0:x0 + valid_width] * 255
            tile_masks[label] = Image.fromarray(mask)
        expected_overlay = render_overlay(Image.fromarray(expected_rgb), tile_masks, alpha=OVERLAY_ALPHA)
        with Image.open(output / "512모델예측오버레이" / path.name) as opened:
            assert np.array_equal(np.asarray(opened), np.asarray(expected_overlay))
    for key, label in enumerate(CLASSES, 1):
        mask_path = output / "모델예측마스크" / label / "unlabelled.png"
        with Image.open(mask_path) as opened:
            assert opened.size == (519, 8)
            mask = np.asarray(opened)
            assert set(np.unique(mask)).issubset({0, 255})
            assert np.array_equal(mask > 0, prediction["masks"][key])
    assert np.any(prediction["masks"][2] & prediction["masks"][3])


def test_model_output_never_reads_ground_truth_or_dataset_json(fixture, monkeypatch):
    forbidden = {fixture["images"] / "unlabelled.json", fixture["images"] / "dataset.json"}
    path_open, builtin_open = Path.open, builtins.open

    def guard(path):
        if isinstance(path, (str, Path)) and Path(path).resolve() in forbidden:
            raise AssertionError("GT or dataset JSON was read by the raw inference pipeline")

    def guarded_path_open(path, *args, **kwargs):
        guard(path)
        return path_open(path, *args, **kwargs)

    def guarded_builtin_open(path, *args, **kwargs):
        guard(path)
        return builtin_open(path, *args, **kwargs)

    monkeypatch.setattr(Path, "open", guarded_path_open)
    monkeypatch.setattr(builtins, "open", guarded_builtin_open)
    assert export(fixture)["rows"] == 3
    record = json.loads((fixture["output"] / "추론계산기록.json").read_text(encoding="utf-8"))
    serialized = json.dumps(record)
    assert "actual_model" in serialized
    assert "unlabelled.json" not in serialized and "dataset.json" not in serialized


def test_empty_predictions_keep_all_tiles_and_header_only_results(fixture):
    fixture["empty"] = True
    result = export(fixture, images=fixture["image"])
    assert result["rows"] == 0 and result["tiles"] == 2
    assert output_rows(fixture["output"]) == []
    assert len(list((fixture["output"] / "512원본타일").glob("*.png"))) == 2
    assert len(list((fixture["output"] / "512모델예측오버레이").glob("*.png"))) == 2
    workbook = load_workbook(fixture["output"] / "damage_results.xlsx", read_only=True)
    try:
        assert list(workbook.active.values) == [tuple(EXPECTED_COLUMNS)]
    finally:
        workbook.close()


def test_existing_output_is_preserved_without_loading_model(fixture):
    fixture["output"].mkdir()
    sentinel = fixture["output"] / "keep.txt"
    sentinel.write_text("existing user result", encoding="utf-8")
    with pytest.raises(FileExistsError):
        export(fixture)
    assert sentinel.read_text(encoding="utf-8") == "existing user result"
    assert not fixture["engine_calls"]


def test_failure_after_saving_first_tile_is_atomic_and_preserves_source(fixture):
    before = source_hashes(fixture["images"])
    parent_contents = set(fixture["output"].parent.iterdir())
    fixture["fail_after_tile"] = True
    with pytest.raises(RuntimeError, match="fixture inference failed"):
        export(fixture)
    assert not fixture["output"].exists()
    assert set(fixture["output"].parent.iterdir()) == parent_contents
    assert source_hashes(fixture["images"]) == before


def test_predictions_with_valid_camera_context_get_physical_values(fixture, monkeypatch):
    enable_metadata(monkeypatch)
    result = export(fixture)
    rows = output_rows(fixture["output"])
    assert result["mapped_rows"] == 3 and result["physical_values_exported"] == 3
    assert len(fixture["context_calls"]) == 1
    for row in rows:
        assert [float(row[key]) for key in WORLD_COLUMNS] == [243001.0, 431001.0, 70.0]
        if row["damage_type"] == "CRC":
            assert float(row["length_m"]) == pytest.approx(float(row["length_px"]) * 0.001)
            assert float(row["width_m"]) == pytest.approx(float(row["width_px"]) * 0.001)
            assert row["area_m2"] == ""
        else:
            assert float(row["area_m2"]) > 0
            assert all(row[key] == "" for key in ("length_px", "width_px", "length_m", "width_m"))
    crc = next(row for row in rows if row["damage_type"] == "CRC")
    # The crossing component's centroid is x=513: the right-hand tile must be
    # first even though the default grid order puts x=0 before x=512.
    assert "x00512" in json.loads(crc["tile_original_paths_json"])[0]
    connection = json.loads((fixture["output"] / "연결정보.json").read_text(encoding="utf-8"))
    tiles = {tile["tile_id"]: tile for tile in connection["tiles"]}
    assert connection["path_base"] == "this_output_folder"
    assert connection["result_source"] == "model_predictions_without_ground_truth"
    assert connection["columns"] == EXPECTED_COLUMNS
    for damage in connection["damages"]:
        for tile_id in damage["tile_ids"]:
            assert damage["damage_id"] in tiles[tile_id]["damage_ids"]


def test_unsafe_gsd_preserves_world_and_pixels_but_withholds_physical_values(fixture, monkeypatch):
    enable_metadata(monkeypatch)
    fixture["ratio"] = 10.001
    result = export(fixture)
    rows = output_rows(fixture["output"])
    assert result["mapped_rows"] == 3 and result["physical_values_exported"] == 0
    assert result["physical_values_withheld"] == 3
    assert all(row[key] != "" for row in rows for key in WORLD_COLUMNS)
    assert all(row[key] == "" for row in rows for key in PHYSICAL_COLUMNS)
    assert float(next(row for row in rows if row["damage_type"] == "CRC")["length_px"]) > 0


def test_multiple_original_photos_load_model_and_mesh_only_once(fixture):
    Image.new("RGB", (519, 8), (0, 0, 0)).save(fixture["images"] / "blank.png")
    result = export(fixture)
    assert result["images"] == 2 and result["tiles"] == 4 and result["rows"] == 3
    assert len(fixture["engine_calls"]) == 1 and len(fixture["mesh_loads"]) == 1
    assert len(list((fixture["output"] / "512원본타일").glob("*.png"))) == 4
    assert len(list((fixture["output"] / "원본사진").glob("*.png"))) == 2


@pytest.mark.parametrize("key,value", [
    ("GpsLatitude", "91"), ("GpsLongitude", "181"),
    ("CalibratedFocalLength", "0"), ("AbsoluteAltitude", "nan"),
    ("CalibratedFocalLength", "not a number"),
])
def test_invalid_metadata_keeps_prediction_but_never_fabricates_coordinates(
    fixture, monkeypatch, key, value,
):
    enable_metadata(monkeypatch)
    metadata = pipeline.read_dji_xmp(fixture["image"])
    metadata[key] = value
    monkeypatch.setattr(pipeline, "read_dji_xmp", lambda path: metadata)
    result = export(fixture)
    rows = output_rows(fixture["output"])
    assert result["rows"] == 3 and result["mapped_rows"] == 0
    assert all(row[name] == "" for row in rows for name in WORLD_COLUMNS + PHYSICAL_COLUMNS)
    assert not fixture["context_calls"]


def test_duplicate_source_stems_are_rejected_before_model_load(fixture):
    Image.new("RGB", (2, 2), (1, 1, 1)).save(fixture["images"] / "unlabelled.JPG")
    with pytest.raises(ValueError, match="unique"):
        export(fixture)
    assert not fixture["output"].exists() and not fixture["engine_calls"]


def test_output_inside_raw_photo_folder_is_rejected(fixture):
    with pytest.raises(ValueError, match="outside"):
        export(fixture, output=fixture["images"] / "generated")
    assert not (fixture["images"] / "generated").exists()
    assert not fixture["engine_calls"]


def test_existing_output_symlink_is_not_followed_or_overwritten(fixture):
    fixture["output"].symlink_to(fixture["images"], target_is_directory=True)
    before = source_hashes(fixture["images"])
    with pytest.raises(FileExistsError):
        export(fixture)
    assert fixture["output"].is_symlink()
    assert source_hashes(fixture["images"]) == before
    assert not fixture["engine_calls"]


def test_wrong_original_mask_shape_cannot_publish_partial_results(fixture, monkeypatch):
    original_predict = pipeline.RawInferenceEngine.predict

    def wrong_shape(self, rgb, *, on_tile=None):
        prediction = original_predict(self, rgb, on_tile=on_tile)
        prediction["class_masks"][1] = prediction["class_masks"][1][:, :-1]
        return prediction

    monkeypatch.setattr(pipeline.RawInferenceEngine, "predict", wrong_shape)
    parent_contents = set(fixture["output"].parent.iterdir())
    with pytest.raises(ValueError, match="original image grid"):
        export(fixture)
    assert not fixture["output"].exists()
    assert set(fixture["output"].parent.iterdir()) == parent_contents


def test_output_created_during_inference_is_not_overwritten(fixture, monkeypatch):
    original_predict = pipeline.RawInferenceEngine.predict

    def concurrently_created(self, rgb, *, on_tile=None):
        prediction = original_predict(self, rgb, on_tile=on_tile)
        fixture["output"].mkdir()
        (fixture["output"] / "keep.txt").write_text("other result", encoding="utf-8")
        return prediction

    monkeypatch.setattr(pipeline.RawInferenceEngine, "predict", concurrently_created)
    with pytest.raises(FileExistsError, match="appeared"):
        export(fixture)
    assert (fixture["output"] / "keep.txt").read_text(encoding="utf-8") == "other result"
    assert list(fixture["output"].iterdir()) == [fixture["output"] / "keep.txt"]
