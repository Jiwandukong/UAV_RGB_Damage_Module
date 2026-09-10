"""The approved demo's 13 result columns plus three image-link columns.

Keep the demo table contract independent of the production reporter: the
published production revision has only the original 13 columns. These CSV,
Excel and JSON helpers preserve the formatting used to create the demo bundle
without importing unpublished production display-tile changes.
"""
from __future__ import annotations

import csv
import json
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

import numpy as np


DAMAGE_NAMES_KO = {"CRC": "균열", "DLM": "박리", "SPL": "박락"}
PRIMARY_COLUMNS = [
    "image",
    "damage_id",
    "damage_type",
    "damage_name_ko",
    "pixel_nodes_json",
    "world_center_x_m",
    "world_center_y_m",
    "world_center_z_m",
    "length_px",
    "length_m",
    "width_px",
    "width_m",
    "area_m2",
    "source_image_path",
    "tile_original_paths_json",
    "tile_overlay_paths_json",
]
DETAIL_COLUMNS = PRIMARY_COLUMNS
SCHEMA_PATH = Path(__file__).resolve().parents[2] / "schemas/damage_results.schema.json"


def write_workbook(
    path: str | Path,
    *,
    rows: list[dict[str, Any]],
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise ImportError("openpyxl is required to write the Excel deliverable") from exc

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    detail_sheet = workbook.active
    detail_sheet.title = "Damage_Details"
    _append_table(detail_sheet, rows, DETAIL_COLUMNS)

    header_fill = PatternFill("solid", fgColor="5595AF")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_cells in sheet.columns:
            letter = column_cells[0].column_letter
            values = [str(cell.value or "") for cell in column_cells[:100]]
            width = min(max(max(map(len, values), default=8) + 2, 10), 60)
            sheet.column_dimensions[letter].width = width
        if sheet.title == "Damage_Details":
            for row in sheet.iter_rows(min_row=2):
                row[4].alignment = Alignment(wrap_text=False)
    workbook.save(path)


def _append_table(sheet: Any, rows: list[dict[str, Any]], columns: Sequence[str]) -> None:
    sheet.append(list(columns))
    for row in rows:
        sheet.append([_excel_value(row.get(column)) for column in columns])


def _excel_value(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(json_ready(value), ensure_ascii=False, separators=(",", ":"))
    if isinstance(value, (np.bool_, bool)):
        return bool(value)
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return float(value)
    return value


def write_csv(
    path: str | Path,
    rows: list[dict[str, Any]],
    fieldnames: Sequence[str] | None = None,
) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    columns = list(fieldnames) if fieldnames is not None else fieldnames_from_rows(rows)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        if not columns:
            return
        writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def fieldnames_from_rows(rows: Sequence[Mapping[str, Any]]) -> list[str]:
    columns: list[str] = []
    for row in rows:
        for key in row:
            if key not in columns:
                columns.append(str(key))
    return columns


def json_ready(value: Any) -> Any:
    if isinstance(value, Mapping):
        return {str(key): json_ready(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_ready(item) for item in value]
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, (np.bool_, bool)):
        return bool(value)
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return float(value)
    if isinstance(value, Path):
        return str(value)
    return value
